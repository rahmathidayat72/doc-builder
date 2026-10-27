import sys
from pathlib import Path

import docx
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.api.export import sanitize_export_filename
from app.core.security import CSRF_HEADER_NAME, SESSION_COOKIE_NAME, session_manager
from app.main import app
from app.services.docx_generator import generate_resume_docx

client = TestClient(app)

SAMPLE_RESUME_PAYLOAD = {
    "id": "77777777-7777-4777-8777-777777777777",
    "title": "Resume Software Architect",
    "templateId": "ats-clean",
    "colorTheme": "#4338ca",
    "fontFamily": "sans",
    "data": {
        "contact": {
            "fullName": "Bambang Pamungkas",
            "email": "bambang@example.id",
            "phone": "081234567890",
            "location": "Jakarta, Indonesia",
            "linkedin": "linkedin.com/in/bambang",
            "website": "bambang.id",
            "websites": ["bambang.id", "github.com/bambang", "bambang.dev"],
        },
        "summary": "Software Architect berpengalaman 8 tahun dalam sistem skala besar.",
        "experiences": [
            {
                "id": "exp-1",
                "jobTitle": "Principal Architect",
                "company": "PT Mega Finansial",
                "location": "Jakarta",
                "startDate": "Jan 2021",
                "endDate": "Sekarang",
                "current": True,
                "description": (
                    "• Memimpin arsitektur sistem core banking microservices.\n"
                    "• Mengoptimalkan throughput sistem 50%."
                ),
            }
        ],
        "education": [
            {
                "id": "edu-1",
                "institution": "Institut Teknologi Bandung",
                "degree": "S2",
                "fieldOfStudy": "Sistem Informasi",
                "startDate": "2018",
                "endDate": "2020",
                "gpa": "3.88",
                "description": "",
            }
        ],
        "skills": [
            {"id": "sk-1", "name": "Python", "category": "Technical"},
            {"id": "sk-2", "name": "TypeScript", "category": "Technical"},
            {"id": "sk-3", "name": "Docker", "category": "Tools"},
        ],
        "languages": [
            {"id": "lang-1", "language": "Bahasa Indonesia", "proficiency": "Native"},
            {"id": "lang-2", "language": "English", "proficiency": "Fluent"},
        ],
        "customSections": [
            {
                "id": "custom-1",
                "sectionTitle": "Sertifikasi",
                "items": [
                    {
                        "id": "custom-item-1",
                        "title": "AWS Certified Developer",
                        "subtitle": "Amazon Web Services",
                        "date": "2025",
                        "description": "Validasi kompetensi pengembangan aplikasi cloud.",
                    }
                ],
            }
        ],
    },
}


def test_generate_resume_docx_creates_valid_document():
    """Unit test service generate_resume_docx menghasilkan berkas docx
    yang valid dan dapat dibaca."""
    docx_io = generate_resume_docx(SAMPLE_RESUME_PAYLOAD)
    assert docx_io is not None
    assert docx_io.getvalue()[:2] == b"PK"  # ZIP/DOCX Magic Bytes

    # Buka dengan python-docx untuk verifikasi integritas
    doc = docx.Document(docx_io)
    text = " ".join([p.text for p in doc.paragraphs])
    assert "BAMBANG PAMUNGKAS" in text
    assert "bambang@example.id" in text
    assert "bambang.id" in text
    assert "github.com/bambang" in text
    assert "bambang.dev" in text
    assert "PT Mega Finansial" in text
    assert "Principal Architect" in text
    assert "Python • TypeScript • Docker" in text
    assert "SERTIFIKASI" in text
    assert "AWS Certified Developer" in text
    assert "Amazon Web Services" in text
    assert "Validasi kompetensi pengembangan aplikasi cloud." in text
    assert "RINGKASAN PROFESIONAL" in text
    assert "PENGALAMAN KERJA" in text
    assert "PENDIDIKAN" in text
    assert "dalam Sistem Informasi" in text


def test_generate_resume_docx_limits_websites_to_three_and_uses_legacy_fallback():
    """Ekspor DOCX memakai maksimal tiga website dan fallback field legacy."""
    four_websites_payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "websites": ["a.dev", "b.dev", "c.dev", "d.dev"],
            },
        },
    }
    docx_io = generate_resume_docx(four_websites_payload)
    text = " ".join([p.text for p in docx.Document(docx_io).paragraphs])
    assert "a.dev" in text
    assert "b.dev" in text
    assert "c.dev" in text
    assert "d.dev" not in text

    legacy_payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "website": "legacy.dev",
                "websites": [""],
            },
        },
    }
    legacy_io = generate_resume_docx(legacy_payload)
    legacy_text = " ".join([p.text for p in docx.Document(legacy_io).paragraphs])
    assert "legacy.dev" in legacy_text


def test_generate_resume_docx_with_english_language():
    """Memverifikasi generate_resume_docx menggunakan label bahasa Inggris
    saat resumeLanguage disetel ke 'en'."""
    en_payload = {**SAMPLE_RESUME_PAYLOAD, "resumeLanguage": "en"}
    docx_io = generate_resume_docx(en_payload)
    assert docx_io is not None
    doc = docx.Document(docx_io)
    text = " ".join([p.text for p in doc.paragraphs])
    assert "PROFESSIONAL SUMMARY" in text
    assert "WORK EXPERIENCE" in text
    assert "EDUCATION" in text
    assert "SKILLS & COMPETENCIES" in text
    assert "LANGUAGE PROFICIENCY" in text
    assert "in Sistem Informasi" in text
    assert "GPA: 3.88" in text


def test_export_docx_endpoint_requires_auth():
    """Memanggil endpoint POST /api/export/docx tanpa session cookie
    harus mengembalikan HTTP 401."""
    resp = client.post("/api/export/docx", json=SAMPLE_RESUME_PAYLOAD)
    assert resp.status_code == 401


def test_export_docx_endpoint_with_auth_and_csrf_returns_file():
    """Memanggil endpoint POST /api/export/docx dengan session & CSRF valid
    mengembalikan berkas docx."""
    # 1. Bootstrap sesi
    boot_resp = client.post(
        "/api/auth/bootstrap", json={"bootstrap_token": session_manager.bootstrap_token}
    )
    assert boot_resp.status_code == 200
    csrf_token = boot_resp.json()["csrf_token"]
    cookie_val = boot_resp.cookies[SESSION_COOKIE_NAME]

    # 2. Panggil export dengan Cookie & CSRF header
    client.cookies.set(SESSION_COOKIE_NAME, cookie_val)
    headers = {CSRF_HEADER_NAME: csrf_token, "Origin": "http://localhost:5173"}

    resp = client.post("/api/export/docx", json=SAMPLE_RESUME_PAYLOAD, headers=headers)
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]
    assert "attachment; filename=" in resp.headers["content-disposition"]
    assert resp.content[:2] == b"PK"


def test_export_pdf_endpoint_returns_422_when_libreoffice_not_installed():
    """Memanggil endpoint POST /api/export/pdf mengembalikan HTTP 422
    jika LibreOffice belum terpasang."""
    # Pastikan sesi aktif
    headers = {CSRF_HEADER_NAME: session_manager.csrf_token, "Origin": "http://localhost:5173"}
    client.cookies.set(SESSION_COOKIE_NAME, session_manager.session_token)

    resp = client.post("/api/export/pdf", json=SAMPLE_RESUME_PAYLOAD, headers=headers)
    # Di Windows environment tanpa LibreOffice, harus mengembalikan 422
    if resp.status_code == 422:
        data = resp.json()
        assert data["error_code"] == "LIBREOFFICE_NOT_FOUND"
        assert "LibreOffice belum terpasang" in data["message"]
    else:
        # Jika LibreOffice terpasang, harus 200 PDF
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"


def test_generate_resume_docx_with_valid_photo():
    """DOCX menyematkan gambar profil saat usePhoto aktif."""
    import base64
    import io

    from PIL import Image

    img = Image.new("RGB", (64, 64), color="green")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    b64_url = f"data:image/jpeg;base64,{base64.b64encode(buf.getvalue()).decode()}"

    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "usePhoto": True,
                "photoUrl": b64_url,
            },
        },
    }
    docx_io = generate_resume_docx(payload)
    doc = docx.Document(docx_io)
    assert len(doc.inline_shapes) == 1


def test_generate_resume_docx_with_use_photo_false_ignores_photo():
    """DOCX tidak menyematkan gambar jika usePhoto bernilai False meskipun photoUrl terisi."""
    import base64
    import io

    from PIL import Image

    img = Image.new("RGB", (64, 64), color="green")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    b64_url = f"data:image/jpeg;base64,{base64.b64encode(buf.getvalue()).decode()}"

    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "usePhoto": False,
                "photoUrl": b64_url,
            },
        },
    }
    docx_io = generate_resume_docx(payload)
    doc = docx.Document(docx_io)
    assert len(doc.inline_shapes) == 0


def test_generate_resume_docx_with_invalid_photo_raises_value_error():
    """DOCX melempar ValueError jika data gambar rusak."""
    import pytest

    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "usePhoto": True,
                "photoUrl": "data:image/jpeg;base64,corruptedbase64",
            },
        },
    }
    with pytest.raises(ValueError, match="Gagal membaca enkoding Base64|tidak valid"):
        generate_resume_docx(payload)


def test_export_docx_endpoint_with_invalid_photo_returns_400():
    """Endpoint export DOCX mengembalikan 400 jika gambar tidak valid."""
    headers = {CSRF_HEADER_NAME: session_manager.csrf_token, "Origin": "http://localhost:5173"}
    client.cookies.set(SESSION_COOKIE_NAME, session_manager.session_token)

    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "usePhoto": True,
                "photoUrl": "data:image/jpeg;base64,not-an-image",
            },
        },
    }
    resp = client.post("/api/export/docx", json=payload, headers=headers)
    assert resp.status_code == 400
    assert "detail" in resp.json()


def test_sanitize_export_filename_menghapus_karakter_berbahaya():
    assert sanitize_export_filename('CV: Budi/2026', "docx") == "CV Budi 2026.docx"
    assert sanitize_export_filename("Resume.pdf", "pdf") == "Resume.pdf"
    assert sanitize_export_filename("../../etc/passwd\n", "pdf") == "etc passwd.pdf"
    assert sanitize_export_filename("   ", "pdf") == "resume.pdf"
    assert sanitize_export_filename(None, "docx") == "resume.docx"


def test_generate_resume_docx_renders_headline_achievements_and_education_description():
    """DOCX merender headline, butir pencapaian terstruktur, dan deskripsi pendidikan."""
    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "contact": {
                **SAMPLE_RESUME_PAYLOAD["data"]["contact"],
                "headline": "Cloud & Backend Architecture Specialist",
            },
            "experiences": [
                {
                    **SAMPLE_RESUME_PAYLOAD["data"]["experiences"][0],
                    "description": "Deskripsi lama tidak boleh dipakai.",
                    "achievements": [
                        "Memimpin migrasi 12 layanan ke Kubernetes.",
                        "Menurunkan latensi API 40%.",
                    ],
                }
            ],
            "education": [
                {
                    **SAMPLE_RESUME_PAYLOAD["data"]["education"][0],
                    "description": "Fokus pada sistem terdistribusi dan keamanan aplikasi.",
                }
            ],
        },
    }

    docx_io = generate_resume_docx(payload)
    doc = docx.Document(docx_io)
    text = " ".join([p.text for p in doc.paragraphs])

    assert "Cloud & Backend Architecture Specialist" in text
    assert "Memimpin migrasi 12 layanan ke Kubernetes." in text
    assert "Menurunkan latensi API 40%." in text
    assert "Deskripsi lama tidak boleh dipakai." not in text
    assert "Fokus pada sistem terdistribusi dan keamanan aplikasi." in text


def test_generate_resume_docx_uses_legacy_description_when_achievements_absent():
    """Draft lama tanpa achievements tetap mengekspor description sebagai butir."""
    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "experiences": [
                {
                    **SAMPLE_RESUME_PAYLOAD["data"]["experiences"][0],
                    "description": "• Butir lama pertama.\n• Butir lama kedua.",
                }
            ],
        },
    }

    docx_io = generate_resume_docx(payload)
    text = " ".join([p.text for p in docx.Document(docx_io).paragraphs])

    assert "Butir lama pertama." in text
    assert "Butir lama kedua." in text


def test_generate_resume_docx_respects_section_order():
    """Urutan sectionOrder diterapkan pada dokumen DOCX."""
    payload = {
        **SAMPLE_RESUME_PAYLOAD,
        "data": {
            **SAMPLE_RESUME_PAYLOAD["data"],
            "sectionOrder": [
                "education",
                "skills",
                "summary",
                "experience",
                "languages",
                "custom",
            ],
        },
    }

    docx_io = generate_resume_docx(payload)
    text = " ".join([p.text for p in docx.Document(docx_io).paragraphs])
    education_index = text.index("PENDIDIKAN")
    skills_index = text.index("KEAHLIAN & KOMPETENSI")
    summary_index = text.index("RINGKASAN PROFESIONAL")

    assert education_index < skills_index < summary_index
