import io
import sys
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.security import CSRF_HEADER_NAME, SESSION_COOKIE_NAME, session_manager
from app.main import app
from app.services.docx_generator import generate_resume_docx
from app.services.resume_parser import extract_contact_info, parse_resume_document

client = TestClient(app)

SAMPLE_PDF_PATH = (
    Path(__file__).resolve().parent.parent.parent
    / "spikes"
    / "pdf2docx-eval"
    / "samples"
    / "01_cv_sample.pdf"
)

SAMPLE_RESUME_PAYLOAD = {
    "title": "Resume Test",
    "data": {
        "contact": {
            "fullName": "Siti Nurhaliza",
            "email": "siti.nurhaliza@example.id",
            "phone": "+6281234567890",
            "location": "Jakarta, Indonesia",
            "linkedin": "",
            "website": "",
        },
        "summary": "Full Stack Developer berpengalaman membangun aplikasi enterprise.",
        "experiences": [
            {
                "id": "1",
                "jobTitle": "Lead Developer",
                "company": "PT Digital Solusi",
                "startDate": "2021",
                "endDate": "Sekarang",
                "current": True,
                "description": (
                    "• Mengelola tim 5 engineer.\n• Menerapkan arsitektur microservices."
                ),
            }
        ],
        "education": [
            {
                "id": "1",
                "institution": "Universitas Indonesia",
                "degree": "S1",
                "fieldOfStudy": "Ilmu Komputer",
            }
        ],
        "skills": [
            {"id": "1", "name": "Python"},
            {"id": "2", "name": "TypeScript"},
            {"id": "3", "name": "FastAPI"},
        ],
    },
}


def test_extract_contact_info_populates_websites_contract():
    """Parser mengisi field legacy dan daftar websites secara konsisten."""
    lines = ["Budi Santoso", "budi@example.com", "budi.dev"]
    contact, _ = extract_contact_info(lines, "\n".join(lines))

    assert contact["website"] == "budi.dev"
    assert contact["websites"] == ["budi.dev"]


def test_extract_contact_info_extracts_professional_headline():
    """Headline di bawah nama dipisahkan dari lokasi dan data kontak."""
    lines = [
        "Budi Santoso",
        "Senior Backend Engineer | Jakarta, Indonesia",
        "budi@example.com | +628123456789",
    ]
    contact, confidence = extract_contact_info(lines, "\n".join(lines))

    assert contact["fullName"] == "Budi Santoso"
    assert contact["headline"] == "Senior Backend Engineer"
    assert contact["location"] == "Jakarta, Indonesia"
    assert confidence["headline"] == 0.55



def test_parse_resume_document_from_pdf():
    """Menguji parsing berkas PDF menghasilkan kontak, ringkasan, dan confidence."""
    assert SAMPLE_PDF_PATH.exists(), f"Berkas sample tidak ditemukan: {SAMPLE_PDF_PATH}"
    pdf_bytes = SAMPLE_PDF_PATH.read_bytes()

    result = parse_resume_document(pdf_bytes, "01_cv_sample.pdf")
    assert result["success"] is True
    data = result["data"]

    # Verifikasi kontak
    assert "budi" in data["contact"]["fullName"].lower()
    assert "budi.sample@example.com" in data["contact"]["email"]
    assert result["confidence"]["email"] >= 0.90
    assert result["confidence"]["overall"] > 0.60


def test_parse_resume_document_from_docx():
    """Menguji parsing berkas DOCX menghasilkan struktur data CV."""
    docx_io = generate_resume_docx(SAMPLE_RESUME_PAYLOAD)
    docx_bytes = docx_io.getvalue()

    result = parse_resume_document(docx_bytes, "resume_sample.docx")
    assert result["success"] is True
    data = result["data"]

    assert "siti nurhaliza" in data["contact"]["fullName"].lower()
    assert "siti.nurhaliza@example.id" in data["contact"]["email"]
    assert len(data["experiences"]) >= 1
    assert len(data["skills"]) >= 1
    assert data["skills"][0]["category"] == "General"


def test_parse_docx_preserves_table_order_and_extra_sections():
    """DOCX bertabel tetap berurutan; bahasa dan section tambahan tidak dibuang."""
    document = Document()
    document.add_paragraph("Rina Wijaya")
    document.add_paragraph("rina@example.com | +628123456789 | Jakarta, Indonesia")
    document.add_paragraph("PENGALAMAN KERJA")
    experience_table = document.add_table(rows=3, cols=1)
    experience_table.cell(0, 0).text = "Senior Engineer"
    experience_table.cell(1, 0).text = "PT Maju Digital"
    experience_table.cell(2, 0).text = "Jan 2021 - Sekarang"
    document.add_paragraph("BAHASA")
    document.add_paragraph("Bahasa Indonesia (Native) | English: Professional")
    document.add_paragraph("SERTIFIKASI")
    document.add_paragraph("AWS Certified Developer")

    output = io.BytesIO()
    document.save(output)
    result = parse_resume_document(output.getvalue(), "resume_table.docx")

    assert len(result["data"]["experiences"]) == 1
    experience = result["data"]["experiences"][0]
    assert experience["jobTitle"] == "Senior Engineer"
    assert experience["company"] == "PT Maju Digital"
    assert experience["current"] is True
    assert [item["language"] for item in result["data"]["languages"]] == [
        "Bahasa Indonesia",
        "English",
    ]
    assert result["data"]["customSections"][0]["sectionTitle"] == "Sertifikasi"
    assert (
        "AWS Certified Developer" in result["data"]["customSections"][0]["items"][0]["description"]
    )


def test_parse_multiline_experience_and_education():
    """Header jabatan/perusahaan sebelum periode dan pendidikan multi-baris terpetakan."""
    document = Document()
    document.add_paragraph("Dimas Pratama")
    document.add_paragraph("dimas@example.com")
    document.add_paragraph("PROFESSIONAL EXPERIENCE")
    document.add_paragraph("Backend Engineer")
    document.add_paragraph("Acme Technology | Bandung")
    document.add_paragraph("Feb 2020 - Dec 2023")
    document.add_paragraph("• Mengurangi latency API 35%")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("Universitas Gadjah Mada")
    document.add_paragraph("2015 - 2019")
    document.add_paragraph("S1 Teknik Informatika | IPK: 3.75 / 4.00")

    output = io.BytesIO()
    document.save(output)
    result = parse_resume_document(output.getvalue(), "resume_multiline.docx")

    experience = result["data"]["experiences"][0]
    assert experience["jobTitle"] == "Backend Engineer"
    assert experience["company"] == "Acme Technology"
    assert experience["location"] == "Bandung"
    assert experience["startDate"].lower() == "feb 2020"
    assert experience["endDate"].lower() == "dec 2023"
    assert "Mengurangi latency API 35%" in experience["description"]

    education = result["data"]["education"][0]
    assert education["institution"] == "Universitas Gadjah Mada"
    assert education["startDate"] == "2015"
    assert education["endDate"] == "2019"
    assert "S1" in education["degree"]
    assert education["gpa"] == "3.75 / 4.00"


def test_parse_resume_with_glyph_bullets_and_graduation_single_year():
    """Parser mengenali bullet kotak/PDF privat, header majemuk, dan tahun kelulusan tunggal."""
    doc = Document()
    doc.add_paragraph("Rasindu Pratama")
    doc.add_paragraph("081234567890 • rasindu@gmail.com • Cileungsi, Bogor")
    doc.add_paragraph("PENGALAMAN KERJA")
    doc.add_paragraph(
        "Staf Administrasi Gudang — PT Sumber Alfaria Trijaya Tbk | Mei 2021 - Sekarang"
    )
    doc.add_paragraph("□ Mengelola data stok barang gudang masuk dan keluar secara sistematis")
    doc.add_paragraph("\uf0b7 Membuat laporan rekonsiliasi stok harian dan bulanan")
    doc.add_paragraph("Kasir — PT Indomarco Prismatama | Jan 2019 - Apr 2021")
    doc.add_paragraph("- Melayani transaksi pembayaran pelanggan dengan cepat dan ramah")
    doc.add_paragraph("PENDIDIKAN")
    doc.add_paragraph("SMA Negeri 1 Cileungsi")
    doc.add_paragraph("Jurusan IPA")
    doc.add_paragraph("Lulus 2015")

    output = io.BytesIO()
    doc.save(output)
    result = parse_resume_document(output.getvalue(), "resume_rasindu.docx")
    assert result["success"] is True

    contact = result["data"]["contact"]
    assert contact["fullName"] == "Rasindu Pratama"
    assert contact["email"] == "rasindu@gmail.com"
    assert contact["phone"] == "081234567890"
    assert "Cileungsi" in contact["location"]
    # Email domain tidak boleh bocor menjadi website
    assert contact["website"] == ""
    assert contact["websites"] == []

    experiences = result["data"]["experiences"]
    assert len(experiences) == 2
    assert experiences[0]["jobTitle"] == "Staf Administrasi Gudang"
    assert experiences[0]["company"] == "PT Sumber Alfaria Trijaya Tbk"
    assert experiences[0]["current"] is True
    assert experiences[0]["startDate"].lower() == "mei 2021"
    assert "Mengelola data stok barang" in experiences[0]["description"]
    assert "Membuat laporan rekonsiliasi" in experiences[0]["description"]
    assert "□" not in experiences[0]["description"]
    assert "\uf0b7" not in experiences[0]["description"]

    assert experiences[1]["jobTitle"] == "Kasir"
    assert experiences[1]["company"] == "PT Indomarco Prismatama"
    assert experiences[1]["startDate"].lower() == "jan 2019"
    assert experiences[1]["endDate"].lower() == "apr 2021"
    assert "Melayani transaksi" in experiences[1]["description"]

    education = result["data"]["education"]
    assert len(education) == 1
    assert education[0]["institution"] == "SMA Negeri 1 Cileungsi"
    assert education[0]["degree"] == "SMA"
    assert education[0]["fieldOfStudy"] == "IPA"
    assert education[0]["startDate"] == ""
    assert education[0]["endDate"] == "2015"


def test_import_resume_endpoint_requires_auth():
    """Endpoint upload /api/import/resume wajib memiliki sesi aktif."""
    files = {"file": ("test.pdf", b"%PDF-1.4 mock content", "application/pdf")}
    resp = client.post("/api/import/resume", files=files)
    assert resp.status_code == 401


def test_import_resume_endpoint_rejects_unsupported_extensions():
    """Endpoint menolak format file selain .pdf dan .docx."""
    # Set autentikasi
    client.cookies.set(SESSION_COOKIE_NAME, session_manager.session_token)
    headers = {CSRF_HEADER_NAME: session_manager.csrf_token, "Origin": "http://localhost:5173"}

    files = {"file": ("malicious.exe", b"binary content", "application/octet-stream")}
    resp = client.post("/api/import/resume", files=files, headers=headers)
    assert resp.status_code == 400
    assert "Format berkas tidak didukung" in resp.json()["detail"]


def test_import_resume_endpoint_success_with_docx():
    """Endpoint berhasil mem-parsing DOCX valid dan mengembalikan data serta confidence."""
    client.cookies.set(SESSION_COOKIE_NAME, session_manager.session_token)
    headers = {CSRF_HEADER_NAME: session_manager.csrf_token, "Origin": "http://localhost:5173"}

    docx_io = generate_resume_docx(SAMPLE_RESUME_PAYLOAD)
    files = {"file": ("test_resume.docx", docx_io.getvalue(), "application/vnd.openxmlformats")}

    resp = client.post("/api/import/resume", files=files, headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert "confidence" in data
    assert data["confidence"]["email"] >= 0.90
    assert "Siti Nurhaliza" in data["data"]["contact"]["fullName"]
