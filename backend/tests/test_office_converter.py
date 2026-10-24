import io
import sys
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.security import CSRF_HEADER_NAME, SESSION_COOKIE_NAME, session_manager
from app.main import app
from app.services.office_converter import (
    OLE_MAGIC_BYTES,
    convert_presentation_to_pdf_sandboxed,
    convert_spreadsheet_to_pdf_sandboxed,
    validate_presentation_magic_bytes,
    validate_spreadsheet_magic_bytes,
)

client = TestClient(app)


def auth_headers() -> dict[str, str]:
    client.cookies.set(SESSION_COOKIE_NAME, session_manager.session_token)
    return {
        CSRF_HEADER_NAME: session_manager.csrf_token,
        "Origin": "http://localhost:5173",
    }


def make_dummy_pptx_bytes() -> bytes:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_dummy_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="Test")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_test_pdf_bytes(page_count: int = 2) -> bytes:
    doc = pymupdf.open()
    for i in range(page_count):
        page = doc.new_page(width=600, height=800)
        page.insert_text((50, 50), f"Page {i + 1} Content")
    data = doc.tobytes()
    doc.close()
    return data


def make_pdf_with_table_bytes() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=600, height=800)
    # Draw table border lines
    for x in [50, 150, 250]:
        page.draw_line((x, 50), (x, 110))
    for y in [50, 80, 110]:
        page.draw_line((50, y), (250, y))
    page.insert_text((60, 70), "Header1")
    page.insert_text((160, 70), "Header2")
    page.insert_text((60, 100), "Value1")
    page.insert_text((160, 100), "=cmd|' /C calc'!A0")  # Formula injection payload
    data = doc.tobytes()
    doc.close()
    return data


# --- VALIDATION TESTS ---


def test_validate_presentation_magic_bytes():
    pptx_bytes = make_dummy_pptx_bytes()
    assert validate_presentation_magic_bytes(pptx_bytes, "demo.pptx") is True
    assert validate_presentation_magic_bytes(OLE_MAGIC_BYTES + b"\x00" * 10, "demo.ppt") is True
    assert validate_presentation_magic_bytes(b"INVALID_HEADER", "demo.pptx") is False
    assert validate_presentation_magic_bytes(pptx_bytes, "demo.txt") is False


def test_validate_spreadsheet_magic_bytes():
    xlsx_bytes = make_dummy_xlsx_bytes()
    assert validate_spreadsheet_magic_bytes(xlsx_bytes, "report.xlsx") is True
    assert validate_spreadsheet_magic_bytes(OLE_MAGIC_BYTES + b"\x00" * 10, "report.xls") is True
    assert validate_spreadsheet_magic_bytes(b"INVALID_HEADER", "report.xlsx") is False
    assert validate_spreadsheet_magic_bytes(xlsx_bytes, "report.doc") is False


def test_convert_office_services_reject_corrupt_bytes():
    with pytest.raises(ValueError, match="tidak valid"):
        convert_presentation_to_pdf_sandboxed(b"CORRUPT", "slide.pptx")

    with pytest.raises(ValueError, match="tidak valid"):
        convert_spreadsheet_to_pdf_sandboxed(b"CORRUPT", "sheet.xlsx")


# --- PRESENTATION TO PDF TESTS ---


def test_convert_presentation_endpoint_requires_auth():
    files = {"file": ("demo.pptx", make_dummy_pptx_bytes(), "application/vnd.openxmlformats")}
    response = client.post("/api/converter/presentation-to-pdf", files=files)
    assert response.status_code in (401, 403)


def test_convert_presentation_endpoint_rejects_wrong_extension():
    files = {"file": ("demo.pdf", b"%PDF-1.4 dummy", "application/pdf")}
    response = client.post(
        "/api/converter/presentation-to-pdf", files=files, headers=auth_headers()
    )
    assert response.status_code == 400
    assert "Harap unggah berkas .pptx atau .ppt" in response.json()["detail"]


def test_convert_presentation_endpoint_libreoffice_handling():
    with patch("app.services.office_converter.detect_libreoffice", return_value=(False, None)):
        files = {
            "file": (
                "demo.pptx",
                make_dummy_pptx_bytes(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        }
        response = client.post(
            "/api/converter/presentation-to-pdf", files=files, headers=auth_headers()
        )
        assert response.status_code == 422
        data = response.json()
        assert data["error_code"] == "LIBREOFFICE_NOT_FOUND"
        assert "LibreOffice belum terpasang" in data["message"]


# --- SPREADSHEET TO PDF TESTS ---


def test_convert_spreadsheet_endpoint_requires_auth():
    files = {"file": ("data.xlsx", make_dummy_xlsx_bytes(), "application/vnd.openxmlformats")}
    response = client.post("/api/converter/spreadsheet-to-pdf", files=files)
    assert response.status_code in (401, 403)


def test_convert_spreadsheet_endpoint_rejects_wrong_extension():
    files = {"file": ("data.docx", b"PK\x03\x04dummy", "application/vnd.openxmlformats")}
    response = client.post(
        "/api/converter/spreadsheet-to-pdf", files=files, headers=auth_headers()
    )
    assert response.status_code == 400
    assert "Harap unggah berkas .xlsx atau .xls" in response.json()["detail"]


def test_convert_spreadsheet_endpoint_libreoffice_handling():
    with patch("app.services.office_converter.detect_libreoffice", return_value=(False, None)):
        files = {
            "file": (
                "data.xlsx",
                make_dummy_xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        response = client.post(
            "/api/converter/spreadsheet-to-pdf", files=files, headers=auth_headers()
        )
        assert response.status_code == 422
        data = response.json()
        assert data["error_code"] == "LIBREOFFICE_NOT_FOUND"
        assert "LibreOffice belum terpasang" in data["message"]


# --- PDF TO PPTX TESTS ---


def test_pdf_to_pptx_endpoint_requires_auth():
    files = {"file": ("doc.pdf", make_test_pdf_bytes(1), "application/pdf")}
    response = client.post("/api/converter/pdf-to-pptx", files=files)
    assert response.status_code in (401, 403)


def test_pdf_to_pptx_endpoint_success():
    pdf_bytes = make_test_pdf_bytes(2)
    files = {"file": ("laporan_tahunan.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/converter/pdf-to-pptx", files=files, headers=auth_headers())
    assert response.status_code == 200
    assert "presentationml.presentation" in response.headers["content-type"]
    assert "laporan_tahunan.pptx" in response.headers["content-disposition"]

    # Validasi bahwa berkas output adalah PPTX valid dengan 2 slide
    prs = Presentation(io.BytesIO(response.content))
    assert len(prs.slides) == 2


# --- PDF TO XLSX TESTS ---


def test_pdf_to_xlsx_endpoint_requires_auth():
    files = {"file": ("doc.pdf", make_test_pdf_bytes(1), "application/pdf")}
    response = client.post("/api/converter/pdf-to-xlsx", files=files)
    assert response.status_code in (401, 403)


def test_pdf_to_xlsx_endpoint_with_tables_and_formula_defense():
    pdf_bytes = make_pdf_with_table_bytes()
    files = {"file": ("rekap_keuangan.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/converter/pdf-to-xlsx", files=files, headers=auth_headers())
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert "rekap_keuangan.xlsx" in response.headers["content-disposition"]

    # Validasi openpyxl membaca workbook dan formula injection dinetralisasi
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Page 1" in wb.sheetnames
    ws = wb["Page 1"]
    # Periksa nilai sel (formula injection harus diawali kutip tunggal)
    val = ws.cell(row=2, column=2).value
    assert val is not None
    assert str(val).startswith("'=")


def test_pdf_to_xlsx_endpoint_returns_422_when_no_tables():
    # PDF hanya teks biasa tanpa tabel
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((50, 50), "Paragraf narasi biasa tanpa tabel apapun.")
    pdf_no_tables = doc.tobytes()
    doc.close()

    files = {"file": ("artikel.pdf", pdf_no_tables, "application/pdf")}
    response = client.post("/api/converter/pdf-to-xlsx", files=files, headers=auth_headers())
    assert response.status_code == 422
    data = response.json()
    assert data["error_code"] == "NO_TABLES_FOUND"
    assert "Tidak ditemukan tabel" in data["message"]


def test_pdf_converters_reject_corrupt_and_empty():
    corrupt_files = {"file": ("bad.pdf", b"%PDF-corrupt", "application/pdf")}
    resp_pptx = client.post(
        "/api/converter/pdf-to-pptx", files=corrupt_files, headers=auth_headers()
    )
    assert resp_pptx.status_code == 422

    resp_xlsx = client.post(
        "/api/converter/pdf-to-xlsx", files=corrupt_files, headers=auth_headers()
    )
    assert resp_xlsx.status_code == 422
