import io
import tempfile
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

from app.api.health import detect_libreoffice
from app.core.sandbox import SandboxConfig, run_sandboxed_subprocess

COLOR_INDIGO = RGBColor(67, 56, 202)
COLOR_DARK = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_GREEN = RGBColor(4, 120, 87)
COLOR_AMBER = RGBColor(180, 83, 9)


def generate_ats_report_docx(analysis: dict[str, Any], resume_title: str = "Resume") -> io.BytesIO:
    """
    Menghasilkan dokumen Microsoft Word (.docx) laporan resmi evaluasi keterbacaan format ATS.
    """
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style standar
    style_normal = doc.styles["Normal"]
    font_normal = style_normal.font
    font_normal.name = "Calibri"
    font_normal.size = Pt(10)
    font_normal.color.rgb = COLOR_DARK

    # 1. Judul Laporan
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("LAPORAN EVALUASI KETERBACAAN ATS")
    run_title.bold = True
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = COLOR_INDIGO
    p_title.paragraph_format.space_after = Pt(2)

    # Subjudul Dokumen
    now_str = datetime.now().strftime("%d %B %Y, %H:%M")
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(
        f"Dokumen: {resume_title}  |  Dihasilkan: {now_str}  |  Doc-builder Offline Suite"
    )
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = COLOR_MUTED
    p_sub.paragraph_format.space_after = Pt(14)

    # 2. Ringkasan Skor Keseluruhan
    score = analysis.get("score", 0)
    grade = analysis.get("grade", "D")
    passed_count = analysis.get("passedChecksCount", 0)
    total_count = analysis.get("totalChecksCount", 0)

    p_score_card = doc.add_paragraph()
    p_score_card.paragraph_format.space_before = Pt(6)
    p_score_card.paragraph_format.space_after = Pt(10)
    run_score_label = p_score_card.add_run(f"SKOR TOTAL: {score} / 100   —   GRADE: {grade}\n")
    run_score_label.bold = True
    run_score_label.font.size = Pt(13)
    run_score_label.font.color.rgb = COLOR_INDIGO

    run_summary_text = p_score_card.add_run(
        f"Status: {passed_count} dari {total_count} indikator keterbacaan lolos uji evaluasi."
    )
    run_summary_text.font.size = Pt(9.5)

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_INDIGO
        return p

    # 3. Rincian 5 Pilar Penilaian
    add_section_header("1. Rincian Skor 5 Pilar Penilaian")
    breakdown = analysis.get("breakdown", {})

    table_breakdown = doc.add_table(rows=1, cols=3)
    table_breakdown.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_breakdown.rows[0].cells
    hdr_cells[0].text = "Pilar Struktur"
    hdr_cells[1].text = "Skor Tercapai"
    hdr_cells[2].text = "Bobot Maksimum"

    # Styling header tabel
    for cell in hdr_cells:
        for p in cell.paragraphs:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(9.5)

    rows_data = [
        ("Kelengkapan Informasi Kontak", f"{breakdown.get('contact', 0)} poin", "25 poin"),
        ("Kelengkapan Struktur Bagian Inti", f"{breakdown.get('structure', 0)} poin", "25 poin"),
        ("Kepadatan & Kualitas Ringkasan", f"{breakdown.get('summary', 0)} poin", "20 poin"),
        ("Pengukuran Dampak & Metrik", f"{breakdown.get('quantified_impact', 0)} poin", "20 poin"),
        ("Integritas Konten & Anti-Stuffing", f"{breakdown.get('integrity', 0)} poin", "10 poin"),
    ]

    for item, actual, max_pts in rows_data:
        row = table_breakdown.add_row().cells
        row[0].text = item
        row[1].text = actual
        row[2].text = max_pts
        for cell in row:
            for p in cell.paragraphs:
                p.runs[0].font.size = Pt(9)

    # 4. Rekomendasi Tindakan Prioritas
    recommendations = analysis.get("recommendations", [])
    if recommendations:
        add_section_header("2. Rekomendasi Tindakan Prioritas")
        for rec in recommendations:
            p_rec = doc.add_paragraph(style="List Bullet")
            p_rec.paragraph_format.space_after = Pt(2)
            run_rec = p_rec.add_run(rec)
            run_rec.font.size = Pt(9)
    else:
        add_section_header("2. Rekomendasi Tindakan Prioritas")
        p_rec = doc.add_paragraph()
        run_rec = p_rec.add_run(
            "Selamat! Dokumen Anda telah memenuhi seluruh kriteria dasar format ATS."
        )
        run_rec.font.size = Pt(9)
        run_rec.font.color.rgb = COLOR_GREEN

    # 5. Checklist Indikator Detail
    checklist = analysis.get("checklist", [])
    if checklist:
        add_section_header("3. Status Checklist Indikator")
        table_chk = doc.add_table(rows=1, cols=4)
        table_chk.alignment = WD_TABLE_ALIGNMENT.CENTER
        chk_hdrs = table_chk.rows[0].cells
        chk_hdrs[0].text = "Kategori"
        chk_hdrs[1].text = "Indikator Pemeriksaan"
        chk_hdrs[2].text = "Status"
        chk_hdrs[3].text = "Catatan Panduan"

        for cell in chk_hdrs:
            for p in cell.paragraphs:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(9)

        for item in checklist:
            row = table_chk.add_row().cells
            row[0].text = item.get("category", "")
            row[1].text = item.get("item", "")
            status_text = "LOLOS" if item.get("passed") else "PERLU CEK"
            row[2].text = status_text
            row[3].text = item.get("tip", "")

            # Set ukuran font
            for idx, cell in enumerate(row):
                for p in cell.paragraphs:
                    p.runs[0].font.size = Pt(8.5)
                    if idx == 2:
                        p.runs[0].bold = True
                        is_pass = item.get("passed", False)
                        p.runs[0].font.color.rgb = COLOR_GREEN if is_pass else COLOR_AMBER

    # 6. Disclaimer Resmi
    p_disc = doc.add_paragraph()
    p_disc.paragraph_format.space_before = Pt(18)
    run_disc_hdr = p_disc.add_run("Catatan Penafian (Disclaimer):\n")
    run_disc_hdr.bold = True
    run_disc_hdr.font.size = Pt(8.5)
    run_disc = p_disc.add_run(
        analysis.get(
            "disclaimer",
            "Skor ATS bersifat heuristik untuk mengevaluasi keterbacaan format mesin "
            "dan bukan merupakan jaminan penerimaan kerja pada proses rekrutmen.",
        )
    )
    run_disc.font.size = Pt(8)
    run_disc.font.color.rgb = COLOR_MUTED

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def generate_ats_report_pdf(
    analysis: dict[str, Any], resume_title: str = "Resume"
) -> tuple[io.BytesIO | None, str | None]:
    """
    Mengonversi laporan evaluasi ATS menjadi berkas PDF via LibreOffice headless.
    """
    lo_available, lo_path = detect_libreoffice()
    if not lo_available or not lo_path:
        return None, "LIBREOFFICE_NOT_FOUND"

    docx_io = generate_ats_report_docx(analysis, resume_title)

    with tempfile.TemporaryDirectory(prefix="docbuilder_report_") as tmpdir:
        work_dir = Path(tmpdir)
        input_docx = work_dir / f"ats_report_{uuid.uuid4().hex[:8]}.docx"
        input_docx.write_bytes(docx_io.getvalue())

        cmd = [
            lo_path,
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--nologo",
            "--convert-to",
            "pdf",
            str(input_docx.name),
            "--outdir",
            str(work_dir),
        ]

        config = SandboxConfig(
            working_dir=work_dir, timeout_sec=60.0, max_memory_mb=500, allow_network=False
        )

        result = run_sandboxed_subprocess(cmd, config)

        expected_pdf = input_docx.with_suffix(".pdf")
        if result.exit_code == 0 and expected_pdf.exists():
            pdf_bytes = expected_pdf.read_bytes()
            out = io.BytesIO(pdf_bytes)
            out.seek(0)
            return out, None

        error_detail = result.error_message or result.stderr or "Gagal membuat PDF laporan ATS."
        return None, f"REPORT_CONVERSION_FAILED: {error_detail}"
