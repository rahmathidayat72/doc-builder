import time
import json
import tracemalloc
from pathlib import Path
import pymupdf
from pdf2docx import Converter
import docx

def evaluate_file(pdf_path: Path, output_docx_path: Path):
    print(f"\n--- Menguji: {pdf_path.name} ---")
    start_time = time.perf_counter()
    tracemalloc.start()
    
    success = False
    error_msg = None
    
    try:
        cv = Converter(str(pdf_path))
        cv.convert(str(output_docx_path), start=0, end=None)
        cv.close()
        success = True
    except Exception as e:
        error_msg = str(e)
        
    current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    duration = time.perf_counter() - start_time
    
    # Verifikasi isi teks & struktur
    source_doc = pymupdf.open(str(pdf_path))
    source_text = "".join([page.get_text() for page in source_doc])
    source_words = set(source_text.lower().split())
    source_doc.close()
    
    docx_text = ""
    table_count = 0
    paragraph_count = 0
    
    if success and output_docx_path.exists():
        doc = docx.Document(str(output_docx_path))
        paragraph_count = len(doc.paragraphs)
        table_count = len(doc.tables)
        
        # Ekstrak seluruh teks dari paragraf dan tabel di docx
        p_texts = [p.text for p in doc.paragraphs if p.text.strip()]
        t_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        t_texts.append(cell.text.strip())
        docx_text = " ".join(p_texts + t_texts)
        docx_words = set(docx_text.lower().split())
        
        # Hitung rasio preservasi kata kunci (vocabulary recall)
        common_words = source_words.intersection(docx_words)
        word_recall = len(common_words) / len(source_words) if source_words else 1.0
    else:
        word_recall = 0.0
        
    result = {
        "file": pdf_path.name,
        "success": success,
        "error": error_msg,
        "duration_sec": round(duration, 3),
        "peak_memory_mb": round(peak / (1024 * 1024), 2),
        "docx_size_kb": round(output_docx_path.stat().st_size / 1024, 2) if output_docx_path.exists() else 0,
        "paragraphs_detected": paragraph_count,
        "tables_detected": table_count,
        "source_word_count": len(source_words),
        "word_recall_ratio": round(word_recall * 100, 2)
    }
    
    print(f"Hasil: Success={success}, Waktu={result['duration_sec']}s, Memori={result['peak_memory_mb']}MB, Recall Kata={result['word_recall_ratio']}%, Tabel={table_count}")
    return result

def main():
    samples_dir = Path(__file__).parent / "samples"
    output_dir = Path(__file__).parent / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    
    pdf_files = sorted(samples_dir.glob("*.pdf"))
    if not pdf_files:
        print("Tidak ada file sample PDF ditemukan.")
        return
        
    results = []
    for pdf_path in pdf_files:
        output_docx = output_dir / f"{pdf_path.stem}.docx"
        res = evaluate_file(pdf_path, output_docx)
        results.append(res)
        
    summary_path = Path(__file__).parent / "benchmark_summary.json"
    with open(summary_path, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)
        
    print(f"\nBenchmark selesai. Ringkasan tersimpan di: {summary_path}")

if __name__ == "__main__":
    main()
